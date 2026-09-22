#!/usr/bin/env python3
"""Normalize TXT/CSV/TSV/JSON/XLSX word lists for crossword generation."""


from __future__ import annotations


import argparse
import csv
import json
import os
import re
import sys
import unicodedata
import zipfile
from pathlib import Path




HEADER_WORDS = {
    "word", "words", "term", "terms", "name", "names", "answer", "answers",
    "display", "board", "entry", "entries", "単語", "人名", "名称", "表示",
    "盤面表記", "クロスワード表記", "答え",
}




def hira_to_kata(text: str) -> str:
    return "".join(chr(ord(ch) + 0x60) if "ぁ" <= ch <= "ゖ" else ch for ch in text)




def has_kana(text: str) -> bool:
    return any("ぁ" <= ch <= "ゖ" or "ァ" <= ch <= "ヿ" or "ㇰ" <= ch <= "ㇿ" for ch in text)




def normalize_word(value: object, profile: str, keep_diacritics: bool) -> tuple[str, str]:
    original = unicodedata.normalize("NFKC", str(value or "").strip())
    active_profile = profile
    if profile == "auto":
        active_profile = "katakana" if has_kana(original) else "latin"


    text = original.upper()
    if active_profile == "katakana":
        text = hira_to_kata(text)
        allowed = []
        removed = []
        for ch in text:
            if (
                "ァ" <= ch <= "ヺ"
                or "ㇰ" <= ch <= "ㇿ"
                or ch in {"ー", "ヽ", "ヾ", "ヿ"}
                or "A" <= ch <= "Z"
                or "0" <= ch <= "9"
            ):
                allowed.append(ch)
            elif not ch.isspace():
                removed.append(ch)
        return "".join(allowed), "".join(removed)


    if not keep_diacritics:
        text = "".join(
            ch for ch in unicodedata.normalize("NFKD", text)
            if not unicodedata.combining(ch)
        )
    allowed = []
    removed = []
    for ch in text:
        if "A" <= ch <= "Z" or "0" <= ch <= "9":
            allowed.append(ch)
        elif keep_diacritics and unicodedata.category(ch).startswith("L"):
            allowed.append(ch)
        elif not ch.isspace():
            removed.append(ch)
    return "".join(allowed), "".join(removed)




def read_text(path: Path) -> list[list[object]]:
    return [[line.strip()] for line in path.read_text(encoding="utf-8-sig").splitlines() if line.strip()]




def read_delimited(path: Path, delimiter: str | None = None) -> list[list[object]]:
    sample = path.read_text(encoding="utf-8-sig")
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(sample[:4096], delimiters=",\t;|").delimiter
        except csv.Error:
            delimiter = ","
    return [row for row in csv.reader(sample.splitlines(), delimiter=delimiter)]




def read_json(path: Path) -> list[list[object]]:
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    if isinstance(data, dict):
        if isinstance(data.get("使用人名"), dict) and isinstance(data["使用人名"].get("values"), list):
            data = data["使用人名"]["values"]
        else:
            data = data.get("entries", data.get("words", data.get("items")))
    if not isinstance(data, list):
        raise ValueError("JSON input must be an array or contain entries/words/items.")
    rows: list[list[object]] = []
    for item in data:
        if isinstance(item, str):
            rows.append([item])
        elif isinstance(item, dict):
            display = item.get("display", item.get("name", item.get("word", item.get("term"))))
            board = item.get("board", item.get("normalized", display))
            rows.append([display, board])
        elif isinstance(item, list):
            rows.append(item)
        else:
            rows.append([item])
    return rows




def read_xlsx(path: Path, sheet_name: str | None) -> list[list[object]]:
    from openpyxl import load_workbook


    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet not found: {sheet_name}")
        sheet = workbook[sheet_name]
    else:
        preferred = ["使用人名", "使用単語", "Words", "Entries", "Sheet1"]
        selected = next((name for name in preferred if name in workbook.sheetnames), workbook.sheetnames[0])
        sheet = workbook[selected]
    return [list(row) for row in sheet.iter_rows(values_only=True)]




def read_rows(path: Path, sheet_name: str | None) -> list[list[object]]:
    suffix = path.suffix.lower()
    if suffix == ".json":
        return read_json(path)
    if suffix == ".txt":
        return read_text(path)
    if suffix == ".tsv":
        return read_delimited(path, "\t")
    if suffix == ".csv":
        return read_delimited(path)
    if suffix in {".xlsx", ".xlsm"} or zipfile.is_zipfile(path):
        return read_xlsx(path, sheet_name)
    if suffix == ".xls":
        raise ValueError("Legacy .xls is not supported. Save it as .xlsx or CSV first.")
    raise ValueError("Supported inputs: .xlsx, .xlsm, .csv, .tsv, .txt, .json")




def looks_like_header(row: list[object]) -> bool:
    tokens = {
        re.sub(r"[\s_\-]+", "", str(value or "").strip().lower())
        for value in row
        if value is not None
    }
    normalized_headers = {re.sub(r"[\s_\-]+", "", word.lower()) for word in HEADER_WORDS}
    return bool(tokens & normalized_headers)




def cell(row: list[object], one_based_column: int | None) -> object | None:
    if one_based_column is None:
        return None
    index = one_based_column - 1
    return row[index] if 0 <= index < len(row) else None




def select_values(
    row: list[object], display_column: int | None, board_column: int | None
) -> tuple[str, str] | None:
    if display_column or board_column:
        display_value = cell(row, display_column or board_column)
        board_value = cell(row, board_column or display_column)
        if display_value is None and board_value is None:
            return None
        return str(display_value or board_value).strip(), str(board_value or display_value).strip()


    strings = [str(value).strip() for value in row if value is not None and str(value).strip()]
    if not strings:
        return None
    nonnumeric = [value for value in strings if not re.fullmatch(r"[\d.]+", value)]
    if not nonnumeric:
        return None
    if len(nonnumeric) >= 2:
        return nonnumeric[0], nonnumeric[1]
    return nonnumeric[0], nonnumeric[0]




def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("input")
    parser.add_argument("output")
    parser.add_argument("--profile", choices=["auto", "katakana", "latin"], default="auto")
    parser.add_argument("--locale", choices=["auto", "ja", "en"], default="auto")
    parser.add_argument("--sheet")
    parser.add_argument("--display-column", type=int)
    parser.add_argument("--board-column", type=int)
    parser.add_argument("--keep-diacritics", action="store_true")
    args = parser.parse_args()


    input_path = Path(args.input).resolve()
    output_path = Path(args.output).resolve()
    rows = read_rows(input_path, args.sheet)
    rows = [row for row in rows if any(value is not None and str(value).strip() for value in row)]
    if rows and looks_like_header(rows[0]):
        rows = rows[1:]


    entries = []
    removed_by_entry = []
    for row_number, row in enumerate(rows, start=1):
        selected = select_values(row, args.display_column, args.board_column)
        if not selected:
            continue
        display, raw_board = selected
        board, removed = normalize_word(raw_board, args.profile, args.keep_diacritics)
        if len(board) < 2:
            raise ValueError(f"Row {row_number}: normalized entry must contain at least two characters: {display!r}")
        entries.append({"display": display, "board": board})
        if removed:
            removed_by_entry.append({"row": row_number, "display": display, "removed": removed})


    if len(entries) < 4:
        raise ValueError("At least four usable entries are required.")


    duplicates: dict[str, list[str]] = {}
    for entry in entries:
        duplicates.setdefault(entry["board"], []).append(entry["display"])
    duplicates = {board: labels for board, labels in duplicates.items() if len(labels) > 1}
    if duplicates:
        detail = "; ".join(f"{board}: {', '.join(labels)}" for board, labels in duplicates.items())
        raise ValueError(f"Duplicate entries after normalization: {detail}")


    detected_locale = args.locale
    if detected_locale == "auto":
        detected_locale = "ja" if any(has_kana(entry["board"]) for entry in entries) else "en"


    payload = {
        "profile": args.profile,
        "locale": detected_locale,
        "sourceFile": os.path.basename(input_path),
        "normalization": {
            "case": "uppercase",
            "hiraganaToKatakana": args.profile in {"auto", "katakana"},
            "diacritics": "preserved" if args.keep_diacritics else "folded-for-latin",
            "removedCharacters": removed_by_entry,
        },
        "entries": entries,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"entries": len(entries), "profile": args.profile, "locale": detected_locale, "warnings": len(removed_by_entry)}, ensure_ascii=False))
    return 0




if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"prepare_words: {exc}", file=sys.stderr)
        raise SystemExit(2)
